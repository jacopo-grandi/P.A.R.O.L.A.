"""Word export of a transcript: speaker side, time, text; unresolved doubtful words highlighted."""

from __future__ import annotations

from io import BytesIO

from ilat.capture.recorder import LABELS
from ilat.schema import Transcript, format_timestamp


def to_docx(t: Transcript, name: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    doc.add_heading(name, level=1)
    doubtful = sum(w.doubtful for s in t.segments for w in s.words)
    if doubtful:
        doc.add_paragraph(f"Parole evidenziate: {doubtful} ancora da verificare (una seconda "
                          "trascrizione dice altro).")  # fmt: skip
    for seg in t.segments:
        p = doc.add_paragraph()
        head = p.add_run(
            f"{LABELS.get(seg.channel or '', '?')} · {format_timestamp(seg.start)[:8]}  "
        )
        head.bold = True
        words = seg.words or []
        if not words:
            p.add_run(seg.text)
            continue
        for k, w in enumerate(words):
            run = p.add_run(w.text + (" " if k < len(words) - 1 else ""))
            if w.doubtful:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
